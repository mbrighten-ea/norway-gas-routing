"""Excel-first report of the frozen diagnostic; no retrieval or optimisation."""
from pathlib import Path
import sys,json,hashlib
import pandas as pd,numpy as np
R=Path.cwd();sys.path.insert(0,str(R/'src'))
from norway_gas_routing.reporting.operational import write_workbook
from norway_gas_routing.reporting.assessment import STYLE,table
from openpyxl import load_workbook
out=R/'outputs/revision_tolerant_2025';raw=next((R/'data/raw/revision_tolerant_2025').iterdir())
freeze=json.loads((out/'provenance/input_freeze.json').read_text())
evaluation=json.loads((out/'provenance/evaluation.json').read_text())
metrics=pd.read_csv(out/'details/metrics.csv');comparison=pd.read_csv(out/'results.csv')
conditional=pd.read_csv(out/'details/conditional_routing.csv')
flows=pd.read_csv(out/'details/frozen_forecasts.csv')
ledger=pd.read_csv(out/'provenance/inputs/reconciliation.csv',parse_dates=['month'])
weights=pd.read_csv(out/'provenance/inputs/weights_2023.csv')
owners=pd.read_csv(out/'provenance/inputs/owners_2023.csv')
overview=[]
for g in ['UK','FR+BE','NL+DE','Common total']:
 p=metrics.loc[metrics.method.eq('Physical156')&metrics.destination.eq(g)].iloc[0]
 e=metrics.loc[metrics.method.eq('EA legacy')&metrics.destination.eq(g)].iloc[0]
 overview.append({'Group':g,'Model MAE bcm/month':p.mae_bcm_month,'EA MAE bcm/month':e.mae_bcm_month,
  'Model WAPE %':p.wape_pct,'EA WAPE %':e.wape_pct,'Model signed bias bcm':p.bias_total_bcm,
  'EA signed bias bcm':e.bias_total_bcm,'n months':int(p.n),'Lower sample error':'Physical routing' if p.mae_bcm_month<e.mae_bcm_month else 'EA legacy'})
overview=pd.DataFrame(overview)
monthly=[];sensitivity=[]
for month in sorted(comparison.target_month.unique()):
 for g in ['UK','FR+BE','NL+DE','Common total']:
  p=comparison.loc[comparison.target_month.eq(month)&comparison.destination.eq(g)].set_index('method')
  monthly.append({'month':pd.Timestamp(month),'Group':g,'Actual bcm':p.at['EA legacy','actual_bcm'],
    'EA forecast bcm':p.at['EA legacy','prediction_bcm'],'Physical forecast bcm':p.at['Physical156','prediction_bcm'],
    'EA signed error bcm':p.at['EA legacy','error_bcm'],'Physical signed error bcm':p.at['Physical156','error_bcm']})
  sensitivity.append(dict(month=pd.Timestamp(month),group=g,reference156_bcm=p.at['Physical156','prediction_bcm'],
   planning143_bcm=p.at['Capacity143','prediction_bcm'],difference_bcm=p.at['Capacity143','prediction_bcm']-p.at['Physical156','prediction_bcm']))
condmetrics=[]
for (method,g),f in conditional.groupby(['method','group']):
 actual=comparison.loc[comparison.method.eq(method)&comparison.destination.eq(g)].actual_bcm
 condmetrics.append(dict(method=method,group=g,n=len(f),conditional_MAE_bcm=f.conditional_routing_error_bcm.abs().mean(),
  conditional_WAPE_pct=100*f.conditional_routing_error_bcm.abs().sum()/actual.abs().sum(),
  conditional_bias_bcm=f.conditional_routing_error_bcm.sum(),
  total_volume_error_component_bcm=f.forecast_total_error_component_bcm.sum()))
condmetrics=pd.DataFrame(condmetrics);condmetrics.to_csv(out/'details/conditional_metrics.csv',index=False)
bindings=pd.read_csv(out/'details/binding_constraints.csv',parse_dates=['month'])
binding_counts=bindings.loc[bindings.scenario.eq('Physical156')&bindings.binding].groupby('asset').size().rename('binding_months').reset_index()
sources=json.loads((raw/'source_metadata.json').read_text())+json.loads((raw/'evaluation_only/source_metadata.json').read_text())
sourceview=pd.DataFrame([dict(source=r.get('name',r.get('key')),sid=r['sid'],role=r['role'],start=r['start'],end=r['end'],
 retrieved_at=r['retrieved_at'],status=r['status'],points=r.get('points'),unit=r.get('metadata',{}).get('unit')) for r in sources])
notes=[
 ('Status','Revision-tolerant historical reconstruction; 31 December 2024 origin, January-December 2025; preliminary 12-month diagnostic, not a verified real-time backtest.'),
 ('Result','Mixed sample ranking: routing lower error for NL+DE/common subset; EA lower error for UK/FR+BE. No claim of overall superiority.'),
 ('Conditional result','After removing common-volume error for scoring, EA has lower destination-share MAE in all three groups. The physical NL+DE raw gain partly reflects cancellation between routing underallocation and total-volume overprediction.'),
 ('Revision policy','Current/revised historical production periods before the origin accepted; retrieval vintages retained. 2025 actuals retrieved only after forecast freeze and used for scoring.'),
 ('Snøhvit','Existing seasonal_two_year method: available month-specific daily-rate mean, one observation if needed, overall historical daily-rate mean only for missing calendar months. This pilot has all 24 observations for 2022/2023, two per calendar month, no imputation.'),
 ('Weights','Requested 2023 calibration window; summed signed residual owners by node, including UNMAPPED. Newer Dec2023-Nov2024 Trym window has eleven missing observations and is not activated. No indicative 2025 weights or 2026-certified 2024 quantities used.'),
 ('Ownership','Existing routes used retrospectively with Core Balder/Urd/Rev decisions. Halten Øst has no 2023 observations and no invented volume; Åsgard equals its native field. Oseberg owns Oseberg/Sør/Øst exactly; Heidrun is explicit once.'),
 ('Missing observations','Tommeliten A has three observed months; Hanz/Trym have no 2023 observations. Source blanks remain blank, not filled. Observed cohort/national identity closes every month; this is aggregate reconciliation, not independent proof of every absent component or future mix.'),
 ('Coverage','0.120104 bcm historical residual production remains UNMAPPED (0.620804% of residual denominator). It is unrouted in forecast. External Trym has no observed calibration numerator; no positive external packet is fabricated. Future external/new-field mix remains uncertain.'),
 ('Capacity','156 mcm/day Kollsnes declared planning ceiling, no assumed maintenance; 143 predeclared capacity bound. Neither is claimed verified as the December-2024 rating. Other existing engineering envelopes unchanged. No 2025 outage data used.'),
 ('UK comparison','Gross model UK minus explicit Martin Linge forecast and the frozen supported non-Gassco UK residual cohort, once. UK_DIRECT has no capacity limit and no disposal. Raw gross UK remains on Physical flows.'),
 ('Common total','Sum of compatible adjusted UK + FR/BE + NL/DE. Excludes DK/PL; not total Norwegian exports. Wider Gassco total remains unscored because boundary/actual-segment contracts do not permit ranking.'),
 ('Conditional diagnostics','Predicted group shares rescaled to observed common total only after freeze for error decomposition. This removes common-volume error for scoring; it is not another issued forecast or an actual-supply LP solve.'),
 ('Physical limits','NL/DE jointly has corridor bounds; individual DE/NL and DK/PL are allocation proxies. Destination choices also depend on unchanged 2023 preferences, not commercial demand or prices.'),
 ('No tuning','All inputs, planning capacities and two LP forecasts froze before target actual retrieval. No parameter changes after scoring. January-September 2026 actuals remain locked.')]
frames={'Summary':overview,'Monthly comparison':pd.DataFrame(monthly),'Conditional routing':condmetrics,
 'Capacity sensitivity':pd.DataFrame(sensitivity),'Production errors':pd.read_csv(out/'details/production_evaluation.csv',parse_dates=['month']),
 'Accounting':ledger,'Source weights':weights,'Owner coverage':owners,'Snøhvit profile':pd.read_csv(out/'provenance/inputs/seasonal_profile.csv'),
 'Physical flows':flows,'Binding corridors':bindings,'Input sources':sourceview,'Read me':pd.DataFrame(notes,columns=['item','meaning'])}
write_workbook(out/'results.xlsx',frames,dict(status='REVISION_TOLERANT_DIAGNOSTIC',analysis_type='Revision-tolerant 2025 forecast comparison, inputs frozen before scoring'))
book=load_workbook(out/'results.xlsx')
for sheet in book:
 for cell in sheet[3]:
  if cell.value and len(str(cell.value))>22:sheet.row_dimensions[3].height=62
 for row in sheet.iter_rows(min_row=4):
  for cell in row:
   label=str(sheet.cell(3,cell.column).value)
   if isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):
    cell.number_format='0' if label in {'n months','n','observations','calendar_month','observed_months','points'} else '0.00"%"' if ('WAPE' in label or label.endswith('_pct')) else '0.00%' if label=='weight' else '0.000'
book.save(out/'results.xlsx');book.close()
# Static scientific comparison figure, using the same frozen scored table.
import matplotlib.pyplot as plt
fig,axes=plt.subplots(2,2,figsize=(12,7))
for ax,g in zip(axes.flat,['UK','FR+BE','NL+DE','Common total']):
 f=pd.DataFrame(monthly);f=f.loc[f.Group.eq(g)].sort_values('month')
 for name,color,style in [('Actual bcm','#202d38','-'),('EA forecast bcm','#939ba5','--'),('Physical forecast bcm','#176b93','-')]:
  ax.plot(f.month,f[name],label=name.replace(' bcm',''),color=color,linestyle=style,linewidth=1.8)
 ax.set_title(g);ax.set_ylabel('bcm/month');ax.grid(alpha=.18);ax.tick_params(axis='x',rotation=30)
axes[0,0].legend(fontsize=8)
fig.suptitle('2025 revision-tolerant reconstruction | December 2024 forecast origin',fontsize=13)
fig.tight_layout();fig.savefig(out/'details/comparison.png',dpi=150);plt.close(fig)
html='<h1>First 2025 accuracy diagnostic</h1><p><strong>Revision-tolerant historical reconstruction, not a verified real-time backtest.</strong> Twelve independent months; no tuning. <a href="results.xlsx">Open the Excel comparison</a>.</p>'+table(overview)
html+='<p>Mixed ranking: physical routing has lower sample errors for NL+DE and the common subset; EA is better for UK and FR+BE. Planning-capacity sensitivity and source-mix assumptions limit any overall conclusion.</p><img src="details/comparison.png" style="max-width:100%">'
html+='<h2>Volume and routing contributions</h2><p>Both use the same archived EA national forecast. Its 2025 bias is +2.840609 bcm. Seasonal LNG is +0.674669 bcm versus revised actual net-saleable LNG, yielding +2.165940 bcm pipeline-production bias. Common subset export bias is +2.847400 bcm for routing and +4.139609 bcm for EA; their boundaries exclude DK/PL and non-Gassco UK packets.</p>'+table(condmetrics)
html+='<p><strong>EA has lower conditional destination-share MAE in all three groups.</strong> NL+DE conditional MAE is 0.240837 bcm/month for the physical model versus 0.139335 for EA. The physical raw-volume advantage partly arises from error cancellation: its conditional NL+DE underallocation (-2.890048 bcm) is offset by +1.262403 bcm from overprediction of the common total. This is not evidence of superior destination-routing skill.</p>'
html+='<h2>Capacity sensitivity, declared before scoring</h2>'+table(metrics.loc[metrics.method.eq('Capacity143'),['destination','n','mae_bcm_month','wape_pct','bias_total_bcm']])
html+='<p>The 143 planning bound reduces exports by 1.099932 bcm and increases disposal by the same amount. NL+DE flows and scores do not change; UK and FR+BE errors remain above EA. This is sensitivity, not a capacity selection based on the scores.</p>'
html+='<h2>Binding corridors and source uncertainty</h2>'+table(binding_counts)+table(weights[['node','annual_bcm','weight']])
html+='<p>Zero slack on an individual parallel arc such as Zeepipe IIB is not independent proof of an economically material bottleneck. These limits are physical/planning assumptions; allocation among feasible exits also follows the fixed historical preferences.</p>'
html+='<p>Residual denominator: 19.346535 bcm; UK-direct weight 34.911735%. Unknown production stays unmapped: 0.120104 bcm, 0.620804% of historical residual. Core routes do not imply field-specific future forecasts: new projects/external production can differ from this historical mix. Balder and Urd each explicitly return twelve zero observations for 2023, Rev returns twelve observations totaling 0.001760 bcm; no overlapping addends were created. Halten Øst has no observations in this calibration year and no volume was added to Åsgard.</p>'
html+='<h2>Assumptions and safeguards</h2>'+table(pd.DataFrame(notes,columns=['item','meaning']))
html+='<p><a href="results.csv">Monthly score table</a> | <a href="provenance/input_freeze.json">Input freeze</a> | <a href="provenance/evaluation.json">Evaluation-only policy</a> | <a href="provenance/manifest.json">Hashes and source captures</a>.</p>'
(out/'report.html').write_text('<!doctype html><html><meta charset="utf-8"><title>2025 accuracy comparison</title><style>'+STYLE+'</style><body>'+html+'</body></html>',encoding='utf-8')
code=out/'provenance/code';code.mkdir(exist_ok=True)
for name in ['capture_inputs.py','prepare_run.py','evaluate.py','report.py']:
 (code/name).write_bytes((R/'.scratch/runs/20261014-01_revision-tolerant-2025'/name).read_bytes())
(code/'seasonal_snohvit.py').write_bytes((R/'src/norway_gas_routing/seasonal_snohvit.py').read_bytes())
files=[p for base in [out,raw] for p in base.rglob('*') if p.is_file() and p.name!='manifest.json']
manifest=dict(base_revision='3313de7e022c77e5197f6687558686966b0dc916',analysis_type='revision-tolerant historical reconstruction',
 forecast_origin=freeze['origin'],source_capture=raw.relative_to(R).as_posix(),operational_model_changed=False,
 protected_actuals_queried=False,parameter_adjustments_after_scoring=False,ordinary_forecast_default='Unchanged SID-backed operational inputs',
 files_sha256={p.relative_to(R).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in files})
(out/'provenance/manifest.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
print(overview.to_string(index=False));print('CONDITIONAL',condmetrics.to_string(index=False));print('BINDING',binding_counts.to_string(index=False))
