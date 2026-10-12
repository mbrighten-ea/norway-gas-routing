"""Audit routing source lineage, using formulas/metadata only; never retrieve data.

Only headers and January/December 2024 formulas are inspected. No workbook
observations (including the reserved 2026 period) are exported or used as inputs.
The resulting ledger is evidence, not an executable source mapping/configuration.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import sys
import unicodedata

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT/'src'))
from norway_gas_routing.config import load_config


def label(value):
    value = str(value).lower().replace('ø', 'o').replace('æ', 'ae')
    return '_'.join(unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode().replace('_', ' ').split())


def audit(output):
    output = (ROOT/output).resolve()
    if not output.is_relative_to(ROOT):
        raise ValueError('Audit output must stay inside this repository.')
    output.mkdir(parents=True, exist_ok=False)
    model, registry = load_config(ROOT/'config/model.yaml', ROOT/'config/data_series.csv')
    discovery_path = ROOT/'evidence/inputs/network_forecast_unverified_discovery.csv'
    discovered = pd.read_csv(discovery_path, usecols=['sid', 'field_label_from_identifier'], keep_default_na=False)
    # Discovery equality is only a search lead, never a verified SID mapping.
    candidates = {}
    for row in discovered.itertuples(index=False):
        candidates.setdefault(label(row.field_label_from_identifier), []).append(row.sid)
    workbook = ROOT/'data/raw/legacy/Norwegian production model.xlsx'
    book = load_workbook(workbook, read_only=True, data_only=False, keep_links=False)
    rows = []
    try:
        sheet = book['Production mcm_d']
        headers = list(sheet.iter_rows(min_row=1, max_row=23))
        history = list(sheet.iter_rows(min_row=432, max_row=443))
        if str(history[0][4].value)[:10] != '2024-01-01' or str(history[-1][4].value)[:10] != '2024-12-01':
            raise ValueError('Historical formula addresses changed; do not guess the date range.')
        for node, entry in model['routing']['nodes'].items():
            for col in entry.get('workbook_columns', []):
                from openpyxl.utils.cell import column_index_from_string
                index = column_index_from_string(col)-1
                field = headers[22][index].value
                first, last = history[0][index].value, history[-1][index].value
                if not all(isinstance(v, str) and v.startswith('=') and 'NPD_raw!' in v for v in [first, last]):
                    raise ValueError(col+': historical lookup formula changed.')
                rows.append(dict(reference=col, field=field, node=node,
                    classification='essential_optimisation', mapping_status='unresolved', verified_key='', verified_sid='',
                    candidate_sid=';'.join(candidates.get(label(field), [])),
                    candidate_status='Discovery only; units/points/identity/EA equivalence unverified',
                    raw_unit='monthly-average mcm/day worksheet reference', frequency='monthly',
                    source_unit='NPD_raw formula declares monthly bcm; candidate API unit unverified',
                    sign='Signed source; no clipping; mapped physical supply must be finite and nonnegative',
                    transformation='Worksheet historical bcm / calendar days * 1000; old adapter reverses with days / 1000',
                    availability='No verified canonical individual EA SID or independently verified 2024 API contract',
                    dependency='Source-node supply -> LP source balance; national-minus-LNG-minus-mapped residual also changes',
                    identifier_cell=str(headers[1][index].value or ''),
                    january_formula=first, december_formula=last,
                    evidence='Production mcm_d!'+col+'23,'+col+'432,'+col+'443; network_forecast_production_inventory.csv'))
    finally:
        book.close()
    ledger = pd.DataFrame(rows)
    if len(ledger) != 46 or ledger.reference.duplicated().any():
        raise ValueError('Expected exactly 46 distinct active workbook-only components.')
    ledger.to_csv(output/'routing_source_mapping.csv', index=False, encoding='utf-8')
    used = [workbook, discovery_path, ROOT/'config/model.yaml', ROOT/'config/data_series.csv', Path(__file__)]
    manifest = dict(scope='Formula/identifier-only 2024 lineage; no observations, live queries or solver',
        essential=46, derived_or_redundant=0, optional_among_46=0, verified_new_mappings=0,
        per_node=dict(Counter(ledger.node)), discovery_candidates=int(ledger.candidate_sid.ne('').sum()),
        source_sha256={p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in used},
        ledger_sha256=hashlib.sha256((output/'routing_source_mapping.csv').read_bytes()).hexdigest())
    (output/'routing_source_audit.json').write_text(json.dumps(manifest, indent=2)+'\n', encoding='utf-8')
    print(json.dumps({k:v for k,v in manifest.items() if 'sha256' not in k}, ensure_ascii=True))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True, type=Path, help='New repository-local evidence directory; never overwrite')
    audit(parser.parse_args().output)
