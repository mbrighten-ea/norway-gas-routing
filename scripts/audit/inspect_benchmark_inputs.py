"""Small read-only probe of forecast metadata and workbook accounting."""
import json
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from helper_functions_ea import ShoojuTools, check_env

ROOT = Path(__file__).resolve().parents[2]

def main():
    cells_path = ROOT / 'evidence/legacy/legacy_workbook_cells.json'
    if not cells_path.exists():
        cells_path = ROOT / 'evidence/legacy/legacy_poc_cells.json'  # Frozen historical evidence.
    cells = json.loads(cells_path.read_text())
    upload = cells['Norwegian flows model.xlsx']['sheets']['SJ upload']
    for row in upload['rows']:
        if any(k in row for k in ['A13','A14','A15','A16','B25']):
            print(row)
    book = load_workbook(ROOT / 'data/raw/legacy/Norwegian production model.xlsx', read_only=True, data_only=False, keep_links=False)
    sheet = book['Production mcm_d']
    for row in sheet.iter_rows(min_row=1,max_row=1):
        print('PRODUCTION HEADER', {c.coordinate:c.value for c in row if c.value is not None})
    print('TOTAL FORMULAS', {c:sheet[c+'432'].value for c in ['HB','HC','HD','HE','HF','HG','HH','HM','HN','IB','IC','IF','HW','L','K','M']})
    print('EXPORT FORMULAS',[(getattr(c,'coordinate',''),c.value) for row in sheet.iter_rows(min_row=432,max_row=432) for c in row if getattr(c,'column',0)>=236 and c.value is not None])
    book.close()
    check_env()
    helper = ShoojuTools(silence_alerts=True)
    sj = helper.sj
    sid = r'teams\natural_gas\EUR\NO\gassco_UK'
    query = f'sid="{sid}"@asof:2023-12-31T23:59:59'
    frame = sj.get_df(query=query, fields=None, max_points=24, df='2023-01-01',dt='2024-12-01')
    print('FRAME',frame)
    print('META',json.dumps(helper.get_fields_from_sid(query,fields=['*']),default=str))
    for sid in [r'teams\natural_gas\EUR\NO_production_Snohvit_field',r'teams\natural_gas\EUR\daily_NO\NO_maintenance_Kollsnes']:
        q=f'sid="{sid}"@asof:2023-12-31T23:59:59'
        frame=sj.get_df(query=q,fields=['series_id','unit','forecast_start_date'],max_points=400,df='2024-01-01',dt='2024-12-31')
        print('MISSING VINTAGE PROBE',sid,frame.shape,frame.head().to_dict('records'))

if __name__ == '__main__':
    main()
