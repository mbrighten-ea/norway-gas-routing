"""Read the two reference workbooks; never save or alter either workbook."""

import io
import hashlib
import json
import struct
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def workbook_bytes(path):
    data = path.read_bytes()
    if zipfile.is_zipfile(io.BytesIO(data)):
        return data, "Original ZIP readable"
    # This local copy has a missing end record and one zeroed directory name.
    # Retain only directory entries matching intact local names, in memory.
    start = data.find(b"PK\x01\x02")
    if start < 0:
        raise ValueError(f"No recoverable ZIP directory: {path.name}")
    position, records = start, []
    while data[position:position + 4] == b"PK\x01\x02":
        name_size, extra_size, comment_size = struct.unpack_from("<HHH", data, position + 28)
        size = 46 + name_size + extra_size + comment_size
        record = data[position:position + size]
        name = record[46:46 + name_size]
        offset = struct.unpack_from("<L", record, 42)[0]
        local_size = struct.unpack_from("<H", data, offset + 26)[0]
        local_name = data[offset + 30:offset + 30 + local_size]
        if name.strip(b"\0") and name == local_name:
            records.append(record)
        position += size
    directory = b"".join(records)
    end = struct.pack("<4s4H2LH", b"PK\x05\x06", 0, 0, len(records), len(records),
                      len(directory), start, 0)
    recovered = data[:start] + directory + end
    with zipfile.ZipFile(io.BytesIO(recovered)) as archive:
        if archive.testzip() is not None:
            raise ValueError("Recovered workbook fails CRC validation")
    return recovered, f"In-memory ZIP directory recovery; {len(records)} entries pass CRC; original untouched"


def inspect():
    result = {}
    selections = {
        "Norwegian production model.xlsx": ["Production bcm", "Production mcm_d", "tech_cap", "Gassco_raw"],
        "Norwegian flows model.xlsx": ["Gassco_raw", "Forecasts", "Baltic Pipe", "SJ upload"],
    }
    for filename, sheets in selections.items():
        original_path = ROOT / "data/raw/legacy" / filename
        original_hash = hashlib.sha256(original_path.read_bytes()).hexdigest()
        data, status = workbook_bytes(original_path)
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        entry = {"read_status": status, "original_sha256": original_hash, "sheets": {}}
        try:
            for name in sheets:
                sheet = formulas[name]
                rows = []
                # Labels/accounting rows plus headers and selected dated flow rows.
                for row_number, row in enumerate(sheet.iter_rows(), 1):
                    values = [cell.value for cell in row]
                    label = " ".join(str(v) for v in values if isinstance(v, str) and not v.startswith("=")).lower()
                    interesting = row_number <= 25 or any(word in label for word in
                        ["gassco", "excl", "snohvit", "snøhvit", "sleipner", "nybro", "faxe", "fergus",
                         "martin", "heimdal", "oseberg", "alvheim", "total norway", "troll"])
                    if any(str(v)[:4].isdigit() and int(str(v)[:4]) >= 2026 and '-' in str(v)[:10] for v in values[:5]):
                        continue
                    interesting |= any(str(v).startswith(date) for v in values[:5]
                                       for date in ["2023-01-01", "2024-01-01"])
                    if interesting:
                        cells = {c.coordinate: getattr(c.value, "text", c.value) for c in row if c.value is not None}
                        rows.append(cells)
                selected = {coordinate for row in rows for coordinate in row}
                caches = {}
                for row in cached[name].iter_rows():
                    for cell in row:
                        if cell.value is not None and cell.coordinate in selected:
                            caches[cell.coordinate] = cell.value
                entry["sheets"][name] = {"rows": rows, "cached_values": caches}
        finally:
            formulas.close()
            cached.close()
        assert hashlib.sha256(original_path.read_bytes()).hexdigest() == original_hash
        entry["original_unchanged"] = True
        result[filename] = entry
    path = ROOT / "evidence/legacy/legacy_workbook_cells.json"
    path.write_text(json.dumps(result, default=str, indent=2) + "\n", encoding="utf-8")
    print(path.relative_to(ROOT))
    for filename, entry in result.items():
        print(filename, entry["read_status"])
        for name, section in entry["sheets"].items():
            print(name, "selected rows", len(section["rows"]))


if __name__ == "__main__":
    inspect()
