"""Physical destination envelopes on the 2025 stage-1 maximum-export face."""
import argparse
from collections import defaultdict
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import numpy as np
import pandas as pd
import pulp
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from norway_gas_routing.artifacts import new_experiment
from norway_gas_routing.config import load_config
from norway_gas_routing.evaluation import check_dates, load_policy
from norway_gas_routing.routing import DESTINATIONS, available_solver, build_month, finite_nonnegative, network_definition, solve_month, validate_solution
from norway_gas_routing.routing_inputs import prepare_inputs
from run_snohvit_sensitivity import official_snohvit, save


def physical_problem(network, supply, capacities):
    """Same node/asset equations as build_month, without preference variables."""
    sources = {n for n, e in network['nodes'].items() if e['kind'] == 'source'}
    if set(supply) != sources:
        raise ValueError('Supply keys must match physical sources')
    finite_nonnegative(supply.values(), 'supply')
    finite_nonnegative(capacities.values(), 'capacity')
    problem = pulp.LpProblem('physical_routing_envelope', pulp.LpMinimize)
    flows = {a['id']: problem.add_variable('flow_' + a['id'], lowBound=0) for a in network['arcs']}
    destinations = {d: problem.add_variable('destination_' + d, lowBound=0) for d in DESTINATIONS}
    unrouted = problem.add_variable('unrouted_supply', lowBound=0)
    incoming, outgoing, assets = defaultdict(list), defaultdict(list), defaultdict(list)
    for arc in network['arcs']:
        incoming[arc['to']].append(flows[arc['id']])
        outgoing[arc['from']].append(flows[arc['id']])
        if arc['asset']:
            assets[arc['asset']].append(flows[arc['id']])
    if set(capacities) != set(assets):
        raise ValueError('Every constrained asset requires its unchanged capacity')
    for node, entry in network['nodes'].items():
        sink = destinations[node] if entry['kind'] == 'destination' else unrouted if entry['kind'] == 'unrouted' else 0
        problem += float(supply.get(node, 0)) + pulp.lpSum(incoming[node]) == pulp.lpSum(outgoing[node]) + sink, 'balance_' + node
    for asset, variables in assets.items():
        problem += pulp.lpSum(variables) <= float(capacities[asset]), 'capacity_' + asset
    problem.setObjective(unrouted)
    return {'problem': problem, 'flows': flows, 'destinations': destinations, 'unrouted': unrouted}


def solved_values(built, solver):
    stats = built['problem'].solve(solver)
    if stats.status != pulp.LpSolveStatus.Optimal:
        raise RuntimeError('Physical envelope solve failed: ' + stats.status_str)
    return {'flows': {k: float(v.varValue) for k, v in built['flows'].items()},
            'destinations': {k: float(v.varValue) for k, v in built['destinations'].items()},
            'unrouted': float(pulp.value(built['unrouted']))}


def envelopes(network, supply, capacities, solver):
    built = physical_problem(network, supply, capacities)
    stage1 = solved_values(built, solver)
    validate_solution(network, supply, capacities, stage1)
    built['problem'] += built['unrouted'] == stage1['unrouted'], 'fix_stage1_maximum_exports'
    endpoints = {}
    for destination in DESTINATIONS:
        for bound, sign in [('minimum', 1), ('maximum', -1)]:
            built['problem'].setObjective(sign * built['destinations'][destination])
            result = solved_values(built, solver)
            checks = validate_solution(network, supply, capacities, result)
            assert abs(result['unrouted'] - stage1['unrouted']) <= 1e-9
            duals = {c.name.removeprefix('capacity_'): c.pi for c in built['problem'].constraints()
                     if c.name.startswith('capacity_')}
            endpoints[destination, bound] = (result, checks, duals)
    return stage1, endpoints


def residual_fields(model, months):
    """Known cached field contributions; unreconciled remainder stays explicit."""
    ledger = pd.read_csv(ROOT / 'outputs/benchmark/details/unclassified_field_components.csv').drop_duplicates('column')
    mapped = {c for entry in model['routing']['nodes'].values() for c in entry.get('workbook_columns', [])}
    evidence = {row['column']: row for row in json.loads((ROOT / 'evidence/inputs/field_transport_evidence.json').read_text(encoding='utf-8'))}
    unassigned = ledger.loc[~ledger.column.isin(mapped)]
    book = load_workbook(ROOT / 'data/raw/legacy/Norwegian production model.xlsx', read_only=True, data_only=True, keep_links=False)
    rows = []
    try:
        sheet = book['Production mcm_d']
        dated = {i: pd.Period(date, freq='M') for i, (date,) in enumerate(sheet.iter_rows(min_col=5, max_col=5, values_only=True), 1)
                 if isinstance(date, datetime) and pd.Period(date, freq='M') in months}
        for i, values in enumerate(sheet.iter_rows(min_row=min(dated), max_row=max(dated), values_only=True), min(dated)):
            if i not in dated:
                continue
            month = dated[i]
            for entry in unassigned.itertuples():
                value = values[column_index_from_string(entry.column) - 1]
                numeric = isinstance(value, (int, float)) and np.isfinite(value)
                rows.append({'month': month, 'column': entry.column, 'cell': entry.column + str(i),
                    'field': evidence.get(entry.column, {}).get('field', entry.header_evidence),
                    'cached_rate_mcm_d': value if numeric else None,
                    'known_signed_volume_bcm': float(value) * month.days_in_month / 1000 if numeric else None,
                    'cell_status': 'retained finite cached value' if numeric else 'unavailable cached value; not filled',
                    'routing_status': 'unclassified under unchanged mapping', 'original_scope_note': entry.scope,
                    'evidence_url': evidence.get(entry.column, {}).get('url', '')})
    finally:
        book.close()
    frame = pd.DataFrame(rows)
    assert frame.month.nunique() == len(months)
    return frame


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run-dir', type=Path)
    args = parser.parse_args()
    run = args.run_dir.resolve() if args.run_dir else new_experiment(ROOT, 'routing-flexibility')
    if not run.is_relative_to(ROOT / '.scratch/runs') or not run.is_dir() or any(run.iterdir()):
        raise ValueError('A new empty named experiment is required')
    paths = [ROOT / '.git/index']
    for folder in ('src', 'config', 'data/raw', 'outputs', 'archive/20261009-13_snohvit-sensitivity'):
        paths.extend(p for p in (ROOT / folder).rglob('*') if p.is_file())
    protected = {p.relative_to(ROOT).as_posix(): sha(p) for p in paths if p.exists()}
    model, registry = load_config(ROOT / 'config/model.yaml', ROOT / 'config/data_series.csv')
    candidate = deepcopy(model)
    candidate['routing'].update(start='2025-01', end='2025-12')
    primary = ROOT / 'data/raw/series_audit/snohvit_sodir_primary_20261008.html'
    inputs = prepare_inputs(ROOT, candidate, registry, analysis_mode='snohvit_sensitivity',
                            diagnostic_scenario='official-net-production', diagnostic_snohvit_bcm=official_snohvit(primary))
    check_dates(inputs['months'], 'analysis')
    check_dates(inputs['training_observations'].index, 'fit')
    network = network_definition(candidate['routing'])
    solver_name, solver = available_solver('HiGHS')
    existing = pd.read_csv(ROOT / 'archive/20261009-13_snohvit-sensitivity/monthly_seven_countries.csv')
    existing = existing.loc[existing.scenario.eq('official-net-production')].set_index('month')
    fields = residual_fields(model, inputs['months'])
    field_sum = fields.groupby('month').known_signed_volume_bcm.sum(min_count=1)
    rows, disposal, residual, constraints, endpoint_flows = [], [], [], [], []
    all_checks = []
    for month in inputs['months']:
        supply, capacity = inputs['supply'].loc[month].to_dict(), inputs['capacities'].loc[month].to_dict()
        baseline = solve_month(build_month(network, supply, capacity, inputs['preferences'], str(month)), solver)
        all_checks.append(validate_solution(network, supply, capacity, baseline))
        old = existing.loc[str(month)]
        for d in DESTINATIONS:
            assert abs(baseline['destinations'][d] - old[d]) <= 1e-9
        assert abs(baseline['unrouted'] - old.unrouted_bcm) <= 1e-9
        stage1, endpoints = envelopes(network, supply, capacity, solver)
        assert abs(stage1['unrouted'] - baseline['unrouted']) <= 1e-9
        for node, volume in supply.items():
            loss = baseline['flows'][node + '_UNROUTED']
            disposal.append({'month': month, 'source': node, 'supply_bcm': volume, 'exported_from_source_bcm': volume - loss,
                'unrouted_bcm': loss, 'classification': 'UNMAPPED accounting residual' if node == 'UNMAPPED' else 'mapped-source disposal at preference baseline'})
        assert abs(baseline['flows']['UNMAPPED_UNROUTED'] - supply['UNMAPPED']) <= 1e-9
        recon = inputs['reconciliation'].loc[month]
        residual.append({'month': month, 'norway_total_bcm': recon.norway_total_production_bcm,
            'official_snohvit_bcm': recon.snohvit_net_production_bcm, 'pipeline_supply_bcm': sum(supply.values()),
            'mapped_supply_bcm': recon.mapped_source_supply_bcm, 'unmapped_unrouted_bcm': supply['UNMAPPED'],
            'mapped_unrouted_bcm': baseline['unrouted'] - supply['UNMAPPED'], 'total_unrouted_bcm': baseline['unrouted'],
            'known_unclassified_workbook_fields_bcm': field_sum.at[month],
            'unreconciled_accounting_vintage_remainder_bcm': supply['UNMAPPED'] - field_sum.at[month],
            'unavailable_unclassified_cells': int(fields.loc[fields.month.eq(month)].known_signed_volume_bcm.isna().sum())})
        for d in DESTINATIONS:
            minimum = endpoints[d, 'minimum'][0]['destinations'][d]
            maximum = endpoints[d, 'maximum'][0]['destinations'][d]
            base = baseline['destinations'][d]
            assert minimum - 1e-8 <= base <= maximum + 1e-8
            observed = float(inputs['observed'].at[month, d])
            distance = max(minimum - observed, observed - maximum, 0)
            rows.append({'month': month, 'destination': d, 'baseline_bcm': base, 'minimum_bcm': minimum, 'maximum_bcm': maximum,
                'potential_reduction_bcm': max(0, base - minimum), 'potential_increase_bcm': max(0, maximum - base),
                'range_width_bcm': maximum - minimum, 'fixed_total_exports_bcm': sum(baseline['destinations'].values()),
                'observed_proxy_bcm': observed, 'observed_minus_baseline_bcm': observed - base,
                'observed_inside_individual_envelope': minimum - 1e-8 <= observed <= maximum + 1e-8,
                'distance_outside_envelope_bcm': distance,
                'individually_reducible_absolute_proxy_error_bcm': max(0, abs(base - observed) - distance),
                'observation_status': 'negative allocation proxy; not physical export' if observed < 0 else 'revised allocation proxy'})
            for bound in ('minimum', 'maximum'):
                result, checks, duals = endpoints[d, bound]
                all_checks.append(checks)
                for asset, cap in capacity.items():
                    flow = sum(result['flows'][a['id']] for a in network['arcs'] if a['asset'] == asset)
                    pi = duals[asset]
                    constraints.append({'month': month, 'destination': d, 'endpoint': bound, 'asset': asset,
                        'flow_bcm': flow, 'capacity_bcm': cap, 'margin_bcm': cap - flow, 'binding': abs(cap - flow) <= 1e-7,
                        'objective_capacity_dual': pi,
                        'endpoint_capacity_shadow_bcm_per_bcm': (-pi if bound == 'maximum' else pi) if pi is not None else None,
                        'dual_scope': 'local LP sensitivity at fixed original exports; not a price or a finite capacity-change experiment'})
                for a in network['arcs']:
                    endpoint_flows.append({'month': month, 'destination': d, 'endpoint': bound, 'arc': a['id'],
                                          'from': a['from'], 'to': a['to'], 'asset': a['asset'], 'flow_bcm': result['flows'][a['id']]})
    flexibility = pd.DataFrame(rows)
    annual = flexibility.groupby('destination', sort=False)[['baseline_bcm', 'minimum_bcm', 'maximum_bcm',
        'potential_reduction_bcm', 'potential_increase_bcm', 'observed_proxy_bcm', 'distance_outside_envelope_bcm',
        'individually_reducible_absolute_proxy_error_bcm']].sum().reset_index()
    annual['months_observed_inside'] = flexibility.groupby('destination', sort=False).observed_inside_individual_envelope.sum().to_numpy()
    for name, frame in [('monthly_flexibility.csv', flexibility), ('annual_flexibility.csv', annual),
        ('residual_supply_breakdown.csv', pd.DataFrame(residual)), ('source_disposal.csv', pd.DataFrame(disposal)),
        ('residual_field_contributions.csv', fields), ('endpoint_constraints.csv', pd.DataFrame(constraints)),
        ('endpoint_arc_flows.csv', pd.DataFrame(endpoint_flows))]:
        save(frame, run, name)
    residual_annual = pd.DataFrame(residual).select_dtypes('number').sum().to_frame().T
    residual_annual.insert(0, 'period', '2025-01..2025-12')
    save(residual_annual, run, 'annual_residual_summary.csv')
    source_annual = pd.DataFrame(disposal).groupby('source', sort=False)[
        ['supply_bcm', 'exported_from_source_bcm', 'unrouted_bcm']].sum().reset_index()
    save(source_annual, run, 'annual_source_disposal.csv')
    field_annual = fields.groupby(['column', 'field'], sort=False).known_signed_volume_bcm.agg(
        known_signed_volume_bcm=lambda values: values.sum(min_count=1), known_months='count').reset_index()
    save(field_annual, run, 'annual_residual_fields.csv')
    baseline_inputs = prepare_inputs(ROOT, model, registry)
    reference = pd.read_csv(ROOT / 'outputs/reconstruction/results.csv').set_index('month')
    replay = []
    for month in baseline_inputs['months']:
        supply, cap = baseline_inputs['supply'].loc[month].to_dict(), baseline_inputs['capacities'].loc[month].to_dict()
        result = solve_month(build_month(network_definition(model['routing']), supply, cap, baseline_inputs['preferences'], str(month)), solver)
        all_checks.append(validate_solution(network_definition(model['routing']), supply, cap, result))
        old = reference.loc[str(month.start_time.date())]
        for d in DESTINATIONS:
            assert abs(result['destinations'][d] - old[d + '_volume_bcm']) <= 1e-9
        assert abs(result['unrouted'] - old.unrouted_supply_bcm) <= 1e-9
        replay.append({'month': month, **result['destinations'], 'unrouted_bcm': result['unrouted']})
    save(pd.DataFrame(replay), run, 'baseline_2024_replay.csv')
    assert protected == {p: sha(ROOT / p) for p in protected}
    manifest = {'run_id': run.name, 'run_at_utc': datetime.now(timezone.utc).isoformat(),
        'status': 'PHYSICAL_ENVELOPES_DIAGNOSTIC_ONLY', 'period': '2025-01..2025-12',
        'scenario': 'official-net-production provisional revised-input reconstruction', 'policy': load_policy(),
        'holdout_loaded': False, 'holdout_scored': False, 'prices_implemented': False, 'preferences_fitted': False,
        'envelope_objective': 'Individual destination min/max with only physical constraints and fixed stage-1 maximum total exports',
        'preferences_used_in_envelopes': False, 'minimax_used_in_envelopes': False,
        'annual_range_note': 'Sums of independent monthly extrema for ONE destination; country extrema cannot be combined',
        'solver': solver_name, 'interpreter': sys.executable, 'inputs': inputs['provenance'],
        'source_sha256': sha(primary), 'max_mass_balance_residual_bcm': max(c['max_mass_balance_residual_bcm'] for c in all_checks),
        'max_capacity_violation_bcm': max(c['max_capacity_violation_bcm'] for c in all_checks),
        'baseline_2024_matches_within_bcm': 1e-9, '2025_baseline_matches_preserved_scenario': True,
        'protected_sha256': protected, 'preservation_verified': True,
        'code_sha256': {str(p.relative_to(ROOT)): sha(p) for p in (Path(__file__), ROOT / 'src/norway_gas_routing/routing.py',
            ROOT / 'src/norway_gas_routing/routing_inputs.py', ROOT / 'scripts/run/run_snohvit_sensitivity.py')}}
    (run / 'manifest.json').write_text(json.dumps(manifest, indent=2, default=str) + '\n', encoding='utf-8')
    print('Run:', run)
    print(annual.to_string(index=False))
    print(pd.DataFrame(residual).select_dtypes('number').sum().to_string())


if __name__ == '__main__':
    main()
