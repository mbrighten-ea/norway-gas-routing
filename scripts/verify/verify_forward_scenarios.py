"""Verify completed F2.4 packages; no SID retrieval, fitting or forecast rerun."""
import argparse,hashlib,json
from pathlib import Path
import sys
import pandas as pd
import numpy as np
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'src'))
from norway_gas_routing.config import load_config
from norway_gas_routing.routing import network_definition


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('summary',help='Repository-relative completed four-scenario summary')
    parser.add_argument('--write-validation',action='store_true',help='Write validation only for a newly generated package; default is read-only')
    parser.add_argument('--replay',action='store_true',help='Replay retained supply/capacity/preference snapshots, never retrieve fresh SIDs')
    args=parser.parse_args(argv); summary=(ROOT/args.summary).resolve()
    assert summary.is_relative_to(ROOT)
    batch=json.loads((summary/'provenance/manifest.json').read_text())
    assert not batch['failures'] and batch['controlled_sources_identical']
    for file,sha in batch.get('files_sha256',{}).items():
        assert hashlib.sha256((summary/file).read_bytes()).hexdigest()==sha,file
    model,registry=load_config(ROOT/'config/model.yaml',ROOT/'config/data_series.csv')
    contract=json.loads((ROOT/'evidence/inputs/routing_sources_b1/verified_sources.json').read_text())
    lineage=pd.read_csv(ROOT/'evidence/inputs/production_forecast_lineage/all_field_components.csv')
    assert lineage.loc[lineage.registered_forecast_key.eq('heidrun_prod'),'includes_in_total_w_snohvit'].eq(True).all()
    first=json.loads((ROOT/next(iter(batch['outputs'].values()))/'provenance/run_manifest.json').read_text())
    months=pd.period_range(first['settings']['start'],first['settings']['end'],freq='M'); outputs={}; raw_hashes={}
    for label,relative in batch['outputs'].items():
        path=ROOT/relative; manifest=json.loads((path/'provenance/run_manifest.json').read_text())
        assert manifest['status']=='OPTIMAL' and not manifest['holdout_scored']
        assert manifest['settings']['origin']==batch['origin']
        assert manifest['network']==network_definition(model['routing'])
        assert all(not r.get('metadata',{}).get('_mock_sid_response') for r in manifest['inputs'])
        for file,sha in manifest['files_sha256'].items():
            assert hashlib.sha256((path/file).read_bytes()).hexdigest()==sha,file
        def monthly(file):
            frame=pd.read_csv(path/file,index_col='month',parse_dates=True)
            frame.index=frame.index.to_period('M'); assert frame.index.equals(months)
            return frame
        flows=monthly('results.csv'); supply=monthly('provenance/inputs/supply.csv')
        capacities=monthly('provenance/inputs/capacities.csv')
        assert flows.solver_status.eq('OPTIMAL').all()
        if args.replay:
            from norway_gas_routing.operational import optimise
            preferences=json.loads((path/'provenance/inputs/preferences.json').read_text())
            _,replayed,_=optimise(dict(months=months,supply=supply,capacities=capacities,preferences=preferences),manifest['network'],manifest['solver'])
            saved=json.loads((path/'provenance/solver_results.json').read_text())
            for month,result in saved.items():
                assert abs(replayed[month]['unrouted']-result['unrouted'])<1e-7
                assert abs(replayed[month]['minimum_baseline_deviation_bcm']-result['minimum_baseline_deviation_bcm'])<1e-7
                for country,value in result['destinations'].items():
                    assert abs(replayed[month]['destinations'][country]-value)<1e-7
                for arc,value in result['flows'].items():
                    assert abs(replayed[month]['flows'][arc]-value)<1e-7
        assert flows.national_balance_residual_bcm.abs().max()<1e-9
        assert flows.residual_bcm.ge(0).all() and supply.ge(0).all().all()
        owners=pd.read_csv(path/'provenance/frozen_2024_owners.csv')
        assert owners.sid.is_unique and owners.component_id.is_unique
        assert abs(owners.annual_bcm.sum()-126.231133)<1e-9
        assert owners.loc[owners.component_id.eq('AQ'),'role'].eq('explicit').all()
        hanz=owners.loc[owners.component_id.eq('GI')].iloc[0]
        assert hanz.observed_months==10 and hanz.approved_structural_zero_months==2
        weights=pd.read_csv(path/'provenance/frozen_residual_weights.csv')
        assert abs(weights.weight.sum()-1)<1e-9 and weights.weight.ge(0).all()
        assert not weights.component_id.eq('AQ').any()
        assert weights.denominator_bcm.sub(18.785806).abs().max()<1e-9
        def raw(key):
            record=next(r for r in manifest['inputs'] if r['key']==key and r.get('kind')=='forecast')
            assert record['returned_asof']==batch['origin']
            v=pd.read_csv(path/record['raw_file'],index_col='date',parse_dates=True).value
            if v.index.tz is not None: v.index=v.index.tz_convert('UTC').tz_localize(None)
            return v
        sn=raw('snohvit_net_prod'); sn.index=sn.index.to_period('M')
        assert sn.index.equals(months)
        assert np.allclose(flows.snohvit_bcm,sn*months.days_in_month/1000,atol=1e-10,rtol=0)
        fields=pd.DataFrame({key:raw(key).set_axis(months) for key in contract['component_columns']})
        assert np.allclose(flows.residual_bcm,raw('norway_prod').to_numpy()-flows.snohvit_bcm-fields.sum(axis=1),atol=1e-10,rtol=0)
        domestic=fields.heidrun_prod if label in {'B','D'} else pd.Series(0.,index=months)
        assert np.allclose(flows.hypothetical_domestic_bcm,domestic,atol=1e-10,rtol=0)
        reconstruction=pd.DataFrame(0.,index=months,columns=supply.columns)
        for key,component in contract['component_columns'].items():
            node=owners.loc[owners.component_id.eq(component),'node'].iloc[0]
            reconstruction[node]+=fields[key]
        for row in weights.itertuples():
            if row.node!='EXTERNAL_TRYM': reconstruction[row.node]+=flows.residual_bcm*row.weight
        reconstruction.KARSTO_FEED-=domestic
        assert np.allclose(supply,reconstruction,atol=1e-9,rtol=0)
        maintenance=raw('kollsnes_maint_d')
        assert maintenance.index.equals(pd.date_range(months.min().start_time,months.max().end_time.normalize()))
        assert maintenance.between(0,156).all()
        factor=143/156 if label in {'C','D'} else 1.
        expected=((156-maintenance)*factor).groupby(maintenance.index.to_period('M')).sum()/1000
        assert np.allclose(capacities.KOLLSNES_EXPORT,expected,atol=1e-10,rtol=0)
        assert not any(r['key']=='kollsnes_upstream_maint_d' for r in manifest['inputs'])
        raw_hashes[label]={}
        for file in (path/'provenance/raw').glob('*.csv'):
            v=pd.read_csv(file,index_col='date',parse_dates=True)
            dates=pd.DatetimeIndex(pd.to_datetime(v.index,utc=True)).tz_localize(None)
            permitted=((dates>=pd.Timestamp('2023-01-01')) & (dates<=pd.Timestamp('2024-12-31'))) | ((dates>=months.min().start_time) & (dates<=months.max().end_time))
            assert permitted.all(),file
            raw_hashes[label][file.name]=hashlib.sha256(file.read_bytes()).hexdigest()
        country=[d+'_volume_bcm' for d in ['UK','FR','BE','NL','DE','DK','PL']]
        assert np.allclose(flows[country].sum(axis=1),flows.total_exports_bcm,atol=1e-9,rtol=0)
        assert np.allclose(flows.total_exports_bcm+flows.unrouted_supply_bcm,supply.sum(axis=1),atol=1e-7,rtol=0)
        book=load_workbook(path/'results.xlsx',read_only=False,data_only=True)
        try:
            assert book['Monthly flows'].max_row==len(months)+3
            assert len(book['Monthly flows']._charts)==1
            assert book['Monthly flows'].cell(4,1).value.strftime('%Y-%m')==str(months.min())
            assert book['Monthly flows'].cell(len(months)+3,1).value.strftime('%Y-%m')==str(months.max())
            for sheet in book:
                for row in sheet.iter_rows():
                    assert not any(c.data_type=='e' for c in row),sheet.title
            assert {'Accounting','Source nodes','Annual flows','Binding limits','Share comparison'}.issubset(book.sheetnames)
        finally: book.close()
        outputs[label]=dict(flows=flows,supply=supply,capacities=capacities,
            digest=manifest['files_sha256']['results.xlsx'],closure_max_bcm=float(flows.national_balance_residual_bcm.abs().max()))
    assert all(x==raw_hashes['A'] for x in raw_hashes.values())
    for label,value in outputs.items():
        assert np.allclose(value['flows'].residual_bcm,outputs['A']['flows'].residual_bcm,atol=1e-10,rtol=0)
        assert value['capacities'].drop(columns='KOLLSNES_EXPORT').equals(outputs['A']['capacities'].drop(columns='KOLLSNES_EXPORT'))
    a=outputs['A']['flows']; dif=pd.read_csv(summary/'details/scenario_differences.csv')
    largest=dif.loc[dif.measure.str.endswith('_volume_bcm')].copy()
    largest['magnitude']=largest.difference_bcm.abs(); largest=largest.sort_values('magnitude',ascending=False)
    shares=pd.read_csv(ROOT/batch['outputs']['A']/'details/historical_share_comparison.csv')
    ranges=pd.read_csv(ROOT/batch['outputs']['A']/'details/physical_ranges.csv')
    binding=pd.read_csv(ROOT/batch['outputs']['A']/'details/binding_constraints.csv')
    report=dict(verified=True,scenario_status={k:'OPTIMAL' for k in outputs},months_each=len(months),monthly_solves=len(months)*len(outputs),saved_input_replay=args.replay,
        origin=batch['origin'],raw_sources_identical=True,protected_actuals_in_raw_captures=False,
        maximum_national_closure_bcm=max(v['closure_max_bcm'] for v in outputs.values()),
        frozen_denominator_bcm=18.785806,residual_proxy_pipeline_pct=float(a.residual_bcm.sum()/a.pipeline_supply_bcm.sum()*100),
        reference=dict(national_bcm=float(a.national_production_bcm.sum()),LNG_bcm=float(a.snohvit_bcm.sum()),
            external_trym_bcm=float(a.external_trym_bcm.sum()),unmapped_bcm=float(a.unmapped_supply_bcm.sum()),
            exports_bcm=float(a.total_exports_bcm.sum()),forced_disposal_bcm=float(a.unrouted_supply_bcm.sum()),
            assumed_eligible_network_coverage_pct=float(a.mapped_supply_bcm.sum()/a.export_eligible_bcm.sum()*100),
            pipeline_boundary_coverage_pct=float(a.mapped_supply_bcm.sum()/a.pipeline_supply_bcm.sum()*100)),
        scenario_exports_bcm={k:float(v['flows'].total_exports_bcm.sum()) for k,v in outputs.items()},
        scenario_disposal_bcm={k:float(v['flows'].unrouted_supply_bcm.sum()) for k,v in outputs.items()},
        largest_monthly_country_differences=largest.head(8).drop(columns='magnitude').to_dict('records'),
        binding_month_counts=binding.loc[binding.binding].groupby('asset').size().to_dict(),
        absolute_share_deviation_bcm=float(shares.model_minus_share_bcm.abs().sum()),
        share_deviation_country_bcm=shares.groupby('destination').model_minus_share_bcm.sum().to_dict(),
        physical_ranges=ranges.to_dict('records'),workbook_results_sha256={k:v['digest'] for k,v in outputs.items()},
        workbook_checks='Typed chronological 38 monthly rows, chart objects, accounting/source/annual/binding/share sheets, no Excel error cells; existing snapshot exporter',
        limitations=['Frozen residual mix is not independently forecast future field membership','Provisional technical reference and other future envelopes are not certified dry-export or outage guarantees',
            'Heidrun bound is an outside-network eligibility scenario, not an offtake schedule or domestic-demand estimate',
            'Monthly LP does not demonstrate daily dispatch; individual physical range endpoints not jointly attainable',
            'DE/NL are shared-landing-pool allocation proxies; DK/PL corridor outputs are not compatible with signed DK net/Nybro legacy readings'])
    if args.write_validation:
        (summary/'provenance/validation.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:report[k] for k in ['verified','months_each','monthly_solves','maximum_national_closure_bcm','reference','scenario_exports_bcm']},indent=2))


if __name__=='__main__': main()
