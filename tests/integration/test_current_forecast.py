"""Future workflow fixtures are explicit demonstrations, never EA forecasts."""
import hashlib
import json
from unittest.mock import Mock

from openpyxl import load_workbook
import pandas as pd
import pytest
import yaml

from norway_gas_routing import operational as op, forecast_inputs as forecast
from norway_gas_routing import project_data
from norway_gas_routing.evaluation import check_dates, EvaluationError
from test_operational import workspace


@pytest.fixture(autouse=True)
def october_forecast_clock(monkeypatch):
    frozen=pd.Timestamp('2026-10-13T12:00:00Z')
    monkeypatch.setattr(pd.Timestamp,'now',lambda tz=None: frozen.tz_convert(tz) if tz else frozen.tz_localize(None))


def configure(root,source='synthetic_demo',**changes):
    path=root/'config/run.yaml'
    request=yaml.safe_load(path.read_text())
    request.update(mode='current_forecast',start=None,end=None,origin=None,horizon_months=12,
        input_source=source,capacity_scenario='synthetic_fixture' if source=='synthetic_demo' else 'baseline_prospective')
    request.update(changes)
    path.write_text(yaml.safe_dump(request))


def verify_hashes(output,manifest):
    for path,digest in manifest['files_sha256'].items():
        assert hashlib.sha256((output/path).read_bytes()).hexdigest()==digest


def test_rolling_12_month_solve_excel_balances_and_nonoverwrite(workspace,monkeypatch):
    configure(workspace)
    output,code=op.execute_run(workspace)
    assert code==0
    manifest=json.loads((output/'provenance/run_manifest.json').read_text())
    assert manifest['status']=='OPTIMAL' and 'SYNTHETIC' in manifest['analysis_type']
    origin,months=forecast.forecast_window(manifest['settings']['origin'])
    assert len(months)==12 and months.min()==pd.Timestamp.now().to_period('M')+1
    assert manifest['evaluation']['analysis_mode']=='current_forecast'
    assert manifest['evaluation']['holdout_loaded'] is False and manifest['evaluation']['holdout_scored'] is False
    frame=pd.read_csv(output/'results.csv')
    assert len(frame)==12 and frame.forecast_origin.eq(origin).all()
    assert (frame.total_exports_bcm+frame.unrouted_supply_bcm-frame.pipeline_supply_proxy_bcm).abs().max()<1e-7
    checks=pd.read_csv(output/'details/solver_checks.csv')
    assert checks.max_mass_balance_residual_bcm.max()<1e-7 and checks.max_capacity_violation_bcm.max()<1e-7
    reconciliation=pd.read_csv(output/'provenance/inputs/reconciliation.csv')
    assert (reconciliation.national_production_bcm-reconciliation.snohvit_bcm-reconciliation.pipeline_supply_bcm).abs().max()<1e-9
    book=load_workbook(output/'results.xlsx')
    assert len(book.sheetnames)==9 and len(book['Monthly flows']._charts)==1
    assert 'SYNTHETIC' in book['Read me']['A1'].value
    assert book['Comparisons']['E4'].value=='Unavailable'
    assert (output/'report.html').is_file() and (output/'details/monthly_flows.png').is_file()
    verify_hashes(output,manifest)
    first=hashlib.sha256((output/'results.xlsx').read_bytes()).hexdigest()
    monkeypatch.setattr(op,'optimise',lambda *a,**k:pytest.fail('Validation-only constructed an LP'))
    second,code=op.execute_run(workspace,validate_only=True)
    assert code==0 and output!=second
    assert hashlib.sha256((output/'results.xlsx').read_bytes()).hexdigest()==first
    assert json.loads((second/'provenance/solver_results.json').read_text())=={}


@pytest.mark.parametrize('changes,message',[
    ({'horizon_months':True},'integer'),({'horizon_months':0},'integer'),
    ({'horizon_months':60},'December 2030'),({'origin':'2026-05-31'},'historical replay'),
    ({'origin':True},'quoted UTC timestamp'),({'capacity_scenario':[]},'named string'),
    ({'origin':'2030-01-01'},'no later than now'),({'start':'2026-01'},'derive'),
    ({'capacity_scenario':'baseline_prospective'},'never mix')])
def test_bad_current_request_rejected_before_io(workspace,changes,message):
    configure(workspace,**changes)
    with pytest.raises(ValueError,match=message): op.execute_run(workspace)
    assert not (workspace/'outputs').exists()


@pytest.mark.parametrize('endpoint,message',[
    ('2031-01','December 2030'),('2026-10','first complete delivery month'),
    ('2030-13','YYYY-MM'),('2030-12-01','YYYY-MM'),('', 'YYYY-MM'),
    (True,'YYYY-MM'),(203012,'YYYY-MM')])
def test_bad_endpoint_rejected_before_io(workspace,endpoint,message):
    configure(workspace,forecast_end_month=endpoint)
    with pytest.raises(ValueError,match=message): op.execute_run(workspace)
    assert not (workspace/'outputs').exists()


def test_horizon_configuration_defaults_count_and_endpoint_precedence(workspace):
    configure(workspace)
    request,_,_,_=op.load_run_config(workspace,'config/run.yaml')
    assert (request['start'],request['end'],request['horizon_months'])==('2026-11','2027-10',12)
    configure(workspace,horizon_months=36)
    request,_,_,_=op.load_run_config(workspace,'config/run.yaml')
    assert (request['start'],request['end'],request['horizon_months'])==('2026-11','2029-10',36)
    configure(workspace,horizon_months=12,forecast_end_month='2030-12')
    request,model,_,_=op.load_run_config(workspace,'config/run.yaml')
    assert (request['start'],request['end'],request['horizon_months'])==('2026-11','2030-12',50)
    assert request['forecast_end_month']=='2030-12' and model['routing']['end']=='2030-12'
    # Removing both optional settings preserves the original 12-month request.
    path=workspace/'config/run.yaml'; raw=yaml.safe_load(path.read_text())
    raw.pop('horizon_months'); raw.pop('forecast_end_month'); path.write_text(yaml.safe_dump(raw))
    assert op.load_run_config(workspace,'config/run.yaml')[0]['horizon_months']==12


def test_endpoint_cannot_relabel_historical_reconstruction(workspace):
    configure(workspace,mode='historical_reconstruction',capacity_scenario='baseline',
        start='2024-01',end='2024-12',forecast_end_month='2030-12')
    with pytest.raises(ValueError,match='only supported for current_forecast'): op.execute_run(workspace)
    assert not (workspace/'outputs').exists()


def test_complete_month_window_uses_utc_information_origin():
    origin,months=forecast.forecast_window('2026-10-12T23:30:00-02:00',forecast_end_month='2030-12')
    assert origin=='2026-10-13T01:30:00+00:00'
    assert months.equals(pd.period_range('2026-11','2030-12',freq='M'))


def test_synthetic_50_month_solve_reporting_order_units_and_validation_only(workspace,monkeypatch):
    configure(workspace,forecast_end_month='2030-12')
    monkeypatch.setattr(forecast,'connect_shooju',lambda:pytest.fail('Synthetic run queried EA'))
    output,code=op.execute_run(workspace)
    assert code==0
    manifest=json.loads((output/'provenance/run_manifest.json').read_text())
    assert manifest['status']=='OPTIMAL' and 'SYNTHETIC' in manifest['analysis_type']
    assert manifest['unit']=='bcm/month' and manifest['settings']['horizon_months']==50
    assert manifest['settings']['forecast_end_month']=='2030-12'
    assert manifest['operational_forecast_available'] is False
    assert not manifest['evaluation']['holdout_loaded'] and not manifest['evaluation']['holdout_scored']
    expected=pd.period_range('2026-11','2030-12',freq='M').to_timestamp()
    results=pd.read_csv(output/'results.csv')
    assert pd.DatetimeIndex(pd.to_datetime(results.month)).equals(expected)
    assert results.forecast_origin.eq(manifest['settings']['origin']).all()
    assert (results.total_exports_bcm+results.unrouted_supply_bcm-results.pipeline_supply_proxy_bcm).abs().max()<1e-7
    for name in ['supply','capacities','reconciliation']:
        frame=pd.read_csv(output/('provenance/inputs/'+name+'.csv'))
        assert pd.DatetimeIndex(pd.to_datetime(frame.month)).equals(expected)
    comparisons=pd.read_csv(output/'details/comparisons.csv')
    assert comparisons.unit.eq('bcm/month').all()
    assert pd.DatetimeIndex(pd.to_datetime(comparisons.month).drop_duplicates()).equals(expected)
    assert comparisons.ea_actual_bcm.isna().all() and comparisons.ea_legacy_forecast_bcm.isna().all()
    arcs=pd.read_csv(output/'details/network_flows.csv')
    assert pd.DatetimeIndex(pd.to_datetime(arcs.month).drop_duplicates()).equals(expected)
    checks=pd.read_csv(output/'details/solver_checks.csv')
    assert len(checks)==50 and checks.max_mass_balance_residual_bcm.max()<1e-7 and checks.max_capacity_violation_bcm.max()<1e-7
    solver_results=json.loads((output/'provenance/solver_results.json').read_text())
    assert list(solver_results)==expected.to_period('M').astype(str).tolist()
    book=load_workbook(output/'results.xlsx')
    sheet=book['Monthly flows']
    assert len(book.sheetnames)==9 and sheet.max_row==53 and len(sheet._charts)==1
    assert [sheet.cell(r,1).value for r in range(4,54)]==list(expected.to_pydatetime())
    assert sheet.cell(53,9).value==pytest.approx(results.total_exports_bcm.iloc[-1])
    assert sheet['A53'].number_format=='yyyy-mm-dd'
    chart=sheet._charts[0]
    assert chart.series[0].cat.numRef.f=="'Monthly flows'!$A$4:$A$53"
    assert chart.series[0].val.numRef.f=="'Monthly flows'!$B$4:$B$53"
    html=(output/'report.html').read_text(encoding='utf-8')
    svg=(output/'details/monthly_flows.svg').read_text(encoding='utf-8')
    assert '50 monthly deliveries' in html and 'bcm/month' in html and 'SYNTHETIC' in html
    assert '2026-11' in svg and '2030-12' in svg and 'bcm/month' in svg
    metadata=json.loads((output/'provenance/source_metadata.json').read_text())
    assert len(metadata)==1 and metadata[0]['kind']=='synthetic_forecast_fixture'
    assert len(metadata[0]['target_months'])==50
    verify_hashes(output,manifest)
    original_optimise=op.optimise
    monkeypatch.setattr(op,'optimise',lambda *a,**k:pytest.fail('Validation-only constructed an LP'))
    second,code=op.execute_run(workspace,validate_only=True)
    assert code==0 and second!=output
    validation=json.loads((second/'provenance/run_manifest.json').read_text())
    assert validation['status']=='VALIDATED_ONLY' and len(pd.read_csv(second/'results.csv'))==50
    assert pd.read_csv(second/'results.csv').total_exports_bcm.isna().all()
    assert json.loads((second/'provenance/solver_results.json').read_text())=={}
    # A second solve with the same fixture gives exactly the same monthly results.
    monkeypatch.setattr(op,'optimise',original_optimise)
    third,code=op.execute_run(workspace)
    assert code==0
    pd.testing.assert_frame_equal(results,pd.read_csv(third/'results.csv'))
    verify_hashes(output,manifest)


def fake_forecast(contract,change=None):
    def retrieve(key,start,end,**kwargs):
        spec=contract['sources'][key]
        frequency=spec['frequency']
        dates=pd.date_range(start,end,freq='D' if frequency=='daily' else 'MS',tz='UTC')
        response=dict(registry=dict(key=key,sid=spec['sid']),asof=kwargs['origin'],
            query='sid="'+spec['sid']+'"@asof:'+pd.Timestamp(kwargs['origin']).strftime('%Y-%m-%dT%H:%M:%S'),
            values=pd.Series(1.,index=dates),metadata=dict(series_id=spec['sid'],source_file=spec['source_file'],
                unit=spec['metadata_unit'],frequency=frequency,lifecycle_stage='forecast',
                forecast_start_date='2026-10-01',_mock_sid_response=True))
        if change: change(key,response)
        return response
    return retrieve


def test_complete_available_sid_coverage_still_blocks_unapproved_supply(workspace,monkeypatch):
    configure(workspace,'fresh_sid')
    contract=json.loads((workspace/'evidence/inputs/routing_sources_b1/verified_sources.json').read_text())
    monkeypatch.setattr(forecast,'connect_shooju',lambda:object())
    monkeypatch.setattr(forecast,'get_prospective_forecast',fake_forecast(contract))
    monkeypatch.setattr(op,'optimise',lambda *a,**k:pytest.fail('Unapproved partial supply entered solver'))
    output,code=op.execute_run(workspace,validate_only=True)
    assert code==2
    manifest=json.loads((output/'provenance/run_manifest.json').read_text())
    assert 'Complete 2024 production ledger failed' in manifest['blocker']
    assert '41 components' in manifest['blocker'] and 'TRYM' in manifest['blocker']
    assert (output/'provenance/production_population.csv').is_file()
    assert (output/'provenance/provisional_capacity_monthly.csv').is_file()
    inventory=pd.read_csv(output/'input_sources.csv')
    assert inventory.classification.eq('essential_population').sum()==69
    assert inventory.loc[inventory.classification.eq('essential_population'),'start'].eq('2024-01').all()
    assert len(manifest['inputs'])==19 and all(r['sha256'] for r in manifest['inputs'])
    assert not (output/'provenance/inputs/supply.csv').exists()
    assert pd.read_csv(output/'results.csv').total_exports_bcm.isna().all()
    verify_hashes(output,manifest)


@pytest.mark.parametrize('scenario',['provisional_kollsnes_143','provisional_kollsnes_156'])
@pytest.mark.parametrize('validate_only',[False,True])
def test_missing_2030_snohvit_blocks_full_horizon_and_preserves_f22_gates(workspace,monkeypatch,scenario,validate_only):
    configure(workspace,'fresh_sid',forecast_end_month='2030-12',capacity_scenario=scenario)
    contract=json.loads((workspace/'evidence/inputs/routing_sources_b1/verified_sources.json').read_text())
    def change(key,response):
        if key=='snohvit_net_prod': response['values']=response['values'].loc[:'2029-12-31']
        if key=='kollsnes_maint_d': response['values'].loc['2027-09-01':'2027-09-13']=156.
    monkeypatch.setattr(forecast,'connect_shooju',lambda:object())
    monkeypatch.setattr(forecast,'get_prospective_forecast',fake_forecast(contract,change))
    monkeypatch.setattr(op,'optimise',lambda *a,**k:pytest.fail('Incomplete extended forecast reached optimiser'))
    output,code=op.execute_run(workspace,validate_only=validate_only)
    manifest=json.loads((output/'provenance/run_manifest.json').read_text())
    assert code==2 and manifest['status']=='BLOCKED_INPUTS' and 'MOCK' in manifest['analysis_type']
    assert '41 components' in manifest['blocker'] and 'not certified' in manifest['blocker']
    assert 'no frozen weights' in manifest['blocker'] and 'Snøhvit forecast conversion and vintage compatibility' in manifest['blocker']
    assert 'snohvit_net_prod: incomplete/duplicate forecast coverage' in manifest['blocker']
    missing=pd.period_range('2030-01','2030-12',freq='M').astype(str).tolist()
    record=next(r for r in manifest['inputs'] if r['key']=='snohvit_net_prod')
    assert record['expected_points']==50 and record['returned_points']==38
    assert record['missing_delivery_months']==missing
    raw=pd.read_csv(output/record['raw_file'])
    assert len(raw)==38 and pd.to_datetime(raw.date,utc=True).max()==pd.Timestamp('2029-12-01',tz='UTC')
    assert len(manifest['inputs'])==19
    others=[r for r in manifest['inputs'] if r['key'] not in {'snohvit_net_prod','kollsnes_maint_d'}]
    assert all(r['status']=='VERIFIED_FORECAST_RESPONSE' and r['expected_points']==50 for r in others)
    blockers=json.loads((output/'provenance/forecast_blockers.json').read_text())
    assert blockers['source_coverage_failures'][0]['missing_delivery_months']==missing
    assert not (output/'provenance/inputs/supply.csv').exists()
    assert not (output/'provenance/inputs/preferences.json').exists()
    assert json.loads((output/'provenance/solver_results.json').read_text())=={}
    flows=pd.read_csv(output/'results.csv')
    assert len(flows)==50 and flows.total_exports_bcm.isna().all()
    assert (output/'provenance/provisional_capacity_monthly.csv').exists()==(scenario=='provisional_kollsnes_156')
    if scenario=='provisional_kollsnes_143': assert 'No clipping, zero floor or rescaling' in manifest['blocker']
    verify_hashes(output,manifest)


@pytest.mark.parametrize('key',['troll_prod','kollsnes_maint_d'])
def test_late_missing_production_or_maintenance_is_not_extended(workspace,monkeypatch,key):
    configure(workspace,'fresh_sid',forecast_end_month='2030-12')
    contract=json.loads((workspace/'evidence/inputs/routing_sources_b1/verified_sources.json').read_text())
    # Remove a month-start production observation or the last daily plan point.
    removed='2030-07-01' if key=='troll_prod' else '2030-12-31'
    def change(name,response):
        if name==key: response['values']=response['values'].drop(pd.Timestamp(removed,tz='UTC'))
    monkeypatch.setattr(forecast,'connect_shooju',lambda:object())
    monkeypatch.setattr(forecast,'get_prospective_forecast',fake_forecast(contract,change))
    monkeypatch.setattr(op,'optimise',lambda *a,**k:pytest.fail('Missing late input reached optimiser'))
    output,code=op.execute_run(workspace)
    manifest=json.loads((output/'provenance/run_manifest.json').read_text())
    assert code==2 and manifest['status']=='BLOCKED_INPUTS'
    record=next(r for r in manifest['inputs'] if r['key']==key)
    assert record['missing_dates']==[removed] and record['missing_delivery_months']==[removed[:7]]
    assert pd.Timestamp(removed,tz='UTC') not in pd.to_datetime(pd.read_csv(output/record['raw_file']).date,utc=True).tolist()
    if key=='kollsnes_maint_d': assert not (output/'provenance/provisional_capacity_monthly.csv').exists()


@pytest.mark.parametrize('case,message',[('gap','coverage'),('unit','units'),('actual','genuine forecast'),
    ('origin','origin mismatch'),('negative','negative production'),('nan','nonfinite')])
def test_source_failure_keeps_raw_and_blocks_optimisation(workspace,monkeypatch,case,message):
    configure(workspace,'fresh_sid',horizon_months=1)
    contract=json.loads((workspace/'evidence/inputs/routing_sources_b1/verified_sources.json').read_text())
    def change(key,response):
        if key!='troll_prod': return
        if case=='gap': response['values']=response['values'].iloc[:0]
        if case=='unit': response['metadata']['unit']='mcm_d'
        if case=='actual': response['metadata']['lifecycle_stage']='actual'
        if case=='origin': response['asof']='2026-10-01T00:00:00Z'
        if case=='negative': response['values'].iloc[0]=-1.
        if case=='nan': response['values'].iloc[0]=float('nan')
    monkeypatch.setattr(forecast,'connect_shooju',lambda:object())
    monkeypatch.setattr(forecast,'get_prospective_forecast',fake_forecast(contract,change))
    monkeypatch.setattr(op,'optimise',lambda *a,**k:pytest.fail('Bad forecast entered optimiser'))
    output,code=op.execute_run(workspace)
    assert code==2
    manifest=json.loads((output/'provenance/run_manifest.json').read_text())
    assert message in manifest['blocker']
    record=next(r for r in manifest['inputs'] if r['key']=='troll_prod')
    assert record['status']=='REJECTED' and (output/record['raw_file']).is_file()


def test_current_archive_query_is_exact_without_unlocking_actuals():
    client=Mock()
    origin,months=forecast.forecast_window(horizon_months=1)
    registry=pd.DataFrame([dict(key='forecast',sid='exact_sid',query_or_transform_hint='')]).set_index('key')
    client.get_fields_from_sid.return_value=dict(lifecycle_stage='forecast',forecast_start_date='2026-10-01')
    client.sj.get_df.return_value=pd.DataFrame(dict(dt=[months.min().start_time],val=[1.]))
    response=project_data.get_prospective_forecast('forecast',str(months.min().start_time.date()),
        str(months.max().end_time.date()),origin=origin,registry=registry,client=client)
    assert '@asof:' in response['query'] and response['asof']==origin
    assert client.sj.get_df.call_args.kwargs['df']>= '2026-10-01'
    client.get_points_from_sid_into_series.assert_not_called()
    client.sj.get_df.return_value=pd.DataFrame(columns=['dt','val'])
    with pytest.raises(ValueError,match='no latest fallback'):
        project_data.get_prospective_forecast('forecast',str(months.min().start_time.date()),
            str(months.max().end_time.date()),origin=origin,registry=registry,client=client)
    client.get_points_from_sid_into_series.assert_not_called()
    with pytest.raises(EvaluationError): check_dates(['2026-01','2026-09'],'read')


def test_prospective_export_cannot_export_locked_or_undeclared_dates(workspace):
    configure(workspace)
    request,_,_,_=op.load_run_config(workspace,'config/run.yaml')
    frame=pd.DataFrame(dict(month=['2026-01-01'],value=[1.]))
    with pytest.raises(ValueError,match='outside its delivery contract'):
        op.save_current_forecast_table(frame,workspace/'unsafe.csv',request)
    assert not (workspace/'unsafe.csv').exists()
    frame.month=request['start']+'-01'
    with pytest.raises(EvaluationError): op.save_csv(frame,workspace/'historical.csv')
    assert not (workspace/'historical.csv').exists()
    configure(workspace,forecast_end_month='2030-12')
    request,_,_,_=op.load_run_config(workspace,'config/run.yaml')
    request['horizon_months']=12  # resolved provenance must agree with the endpoint
    with pytest.raises(ValueError,match='declared forecast horizon'):
        op.save_current_forecast_table(frame,workspace/'inconsistent.csv',request)
    assert not (workspace/'inconsistent.csv').exists()
