"""Full physical model with manufactured SID responses, never live observations."""
import hashlib
import importlib.util
import json
import shutil

import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook

from norway_gas_routing import operational as op, routing_sources as sources
from norway_gas_routing.config import load_config
from test_fresh_routing_inputs import response_factory


def npd_response(spec, dates):
    response = dict(registry=dict(key=spec['npd_field']+'_net_actual', sid=spec['sid']),
        asof=None, values=pd.Series(.01, index=dates), metadata=dict(
            unit='bcm', frequency='monthly', lifecycle_stage='actual',
            series_id=spec['sid'], row_level='field', country_iso='NO',
            ea_field_name_standardised=spec['npd_field'], ea_production_type='marketed_production',
            economic_property='production', energy_product='natural_gas',
            source='Norwegian Petroleum Directorate',
            source_url='https://factpages.npd.no/en/field/TableView/Production/Saleable',
            source_obj=dict(aspect='gas-saleable', units='bill Sm3', discovery_field_name=spec['npd_field_label']),
            _mock_sid_response=True))
    if spec['npd_field']=='hanz':
        response['values'].loc[dates<pd.Timestamp('2024-04-01')]=0.
    return response


@pytest.fixture
def npd(repo_root):
    spec=sources.source_contract(repo_root)['sources']['sindre_net_actual']
    dates=pd.date_range('2024-01-01','2024-12-01',freq='MS')
    return spec, dates, npd_response(spec,dates)


def test_actual_monthly_volume_without_forecast_metadata(npd):
    spec, dates, response=npd
    assert sources.verify_response('sindre_net_actual',response,spec,dates,'monthly').sum()==pytest.approx(.12)
    response['metadata']['series_id']='NPD\\monthly_by_field\\SINDRE\\Gas - saleable'
    response['metadata']['series_id_name']='marketed_production\\sindre\\gas-saleable'
    assert len(sources.verify_response('sindre_net_actual',response,spec,dates,'monthly'))==12
    response['metadata']['series_id']='NPD\\monthly_by_field\\OTHER\\Gas - saleable'
    with pytest.raises(ValueError,match='SID mismatch'):
        sources.verify_response('sindre_net_actual',response,spec,dates,'monthly')


@pytest.mark.parametrize('case,message',[
    ('field','field identity'),('measure','saleable-gas'),('unit','units/frequency'),
    ('source_units','billion-standard'),('vintage','vintage substitution'),
    ('forecast','historical monthly'),('missing','missing 2024-01-01, 2024-02-01'),
    ('duplicate','duplicate timestamp'),('nonfinite','missing/nonfinite')])
def test_npd_contract_rejects_incompatible_response(npd,case,message):
    spec,dates,response=npd
    if case=='field': response['metadata']['ea_field_name_standardised']='other'
    if case=='measure': response['metadata']['source_obj']['aspect']='gross-gas'
    if case=='unit': response['metadata']['unit']='mcm_d'
    if case=='source_units': response['metadata']['source_obj']['units']='million Sm3'
    if case=='vintage': response['asof']='2023-12-31'
    if case=='forecast': response['metadata']['lifecycle_stage']='forecast'
    if case=='missing': response['values']=response['values'].iloc[2:]
    if case=='duplicate': response['values']=pd.concat([response['values'],response['values'].iloc[:1]])
    if case=='nonfinite': response['values'].iloc[0]=float('nan')
    with pytest.raises(ValueError,match=message):
        sources.verify_response('sindre_net_actual',response,spec,dates,'monthly')


def test_full_network_mock_solve_and_genuine_missing_month_failure(repo_root,tmp_path,monkeypatch):
    for directory in ['config','src/norway_gas_routing','evidence/inputs/routing_sources_b1']:
        shutil.copytree(repo_root/directory,tmp_path/directory)
    (tmp_path/'scripts/run').mkdir(parents=True)
    shutil.copyfile(repo_root/'scripts/run/run_routing.py',tmp_path/'scripts/run/run_routing.py')
    run_path=tmp_path/'config/run.yaml'
    request=yaml.safe_load(run_path.read_text())
    request.update(mode='historical_reconstruction',start='2024-01',end='2024-12',origin=None,capacity_scenario='baseline')
    run_path.write_text(yaml.safe_dump(request))
    contract=sources.source_contract(tmp_path)
    calls=[]
    ea=response_factory(contract,calls)
    missing=False
    def retrieve(key,start,end,**kwargs):
        spec=contract['sources'][key]
        if spec.get('provider_contract')=='npd_historical':
            response=npd_response(spec,pd.date_range(start,end,freq='MS'))
            if missing and key=='hanz_net_actual': response['values']=response['values'].drop(pd.Timestamp('2024-04-01'))
            return response
        response=ea(key,start,end,**kwargs)
        if key in contract['component_columns']: response['values'][:]=.01
        return response
    monkeypatch.setattr('norway_gas_routing.data_functions.connect_shooju',lambda:object())
    monkeypatch.setattr(sources,'get_registered_series',retrieve)
    output,code=op.execute_run(tmp_path)
    assert code==0
    manifest=json.loads((output/'provenance/run_manifest.json').read_text())
    assert manifest['status']=='OPTIMAL' and 'MOCK SID' in manifest['analysis_type']
    assert sum(r['category']=='production_npd' for r in manifest['inputs'])==46
    supply=pd.read_csv(output/'provenance/inputs/supply.csv',index_col=0)
    expected=pd.Series([.61]*3+[.62]*9,index=supply.index)
    assert (supply.drop(columns='UNMAPPED').sum(axis=1)-expected).abs().max()<1e-10
    results=pd.read_csv(output/'results.csv')
    assert (results.total_exports_bcm+results.unrouted_supply_bcm-results.pipeline_supply_proxy_bcm).abs().max()<1e-7
    assert len(load_workbook(output/'results.xlsx').sheetnames)==9
    for file,digest in manifest['files_sha256'].items():
        assert hashlib.sha256((output/file).read_bytes()).hexdigest()==digest
    missing=True
    monkeypatch.setattr(op,'optimise',lambda *a,**k:pytest.fail('Missing Hanz entered solver'))
    failed,code=op.execute_run(tmp_path)
    assert code==2
    failure=json.loads((failed/'provenance/run_manifest.json').read_text())
    assert 'hanz_net_actual' in failure['blocker'] and '2024-04-01' in failure['blocker']
    record=next(r for r in failure['inputs'] if r['key']=='hanz_net_actual')
    assert record['status']=='REJECTED' and (failed/record['raw_file']).is_file()


def test_unmapped_or_duplicate_historical_component(repo_root,monkeypatch,tmp_path):
    shutil.copytree(repo_root/'evidence/inputs/routing_sources_b1',tmp_path/'evidence/inputs/routing_sources_b1')
    model,registry=load_config(repo_root/'config/model.yaml',repo_root/'config/data_series.csv')
    contract=sources.source_contract(tmp_path)
    contract['historical_components'].pop('GI')
    (tmp_path/sources.CONTRACT).write_text(json.dumps(contract))
    requirements,_,_=sources.source_requirements(tmp_path,model,registry)
    assert requirements.loc[requirements.blocked,'reference'].tolist()==['GI']
    contract['historical_components']['GI']=dict(key=contract['historical_components']['FM']['key'],node='UK_DIRECT')
    (tmp_path/sources.CONTRACT).write_text(json.dumps(contract))
    with pytest.raises(ValueError,match='identity mismatch|Duplicate production'):
        sources.source_requirements(tmp_path,model,registry)


@pytest.mark.parametrize('reference,field',[('CV','fram'),('GI','hanz')])
def test_activation_preserves_registry_and_validates_full_configuration(repo_root,tmp_path,monkeypatch,reference,field):
    for directory in ['config','evidence/inputs/routing_sources_b1']:
        shutil.copytree(repo_root/directory,tmp_path/directory)
    registry_path=tmp_path/'config/data_series.csv'
    lines=registry_path.read_bytes().splitlines(keepends=True)
    key=field+'_net_actual'
    original=b''.join(line for line in lines if not line.startswith((key+',').encode()))
    registry_path.write_bytes(original)
    loader=importlib.util.spec_from_file_location('activation',repo_root/'scripts/audit/activate_historical_fields.py')
    module=importlib.util.module_from_spec(loader)
    loader.loader.exec_module(module)
    monkeypatch.setattr(module,'ROOT',tmp_path)
    monkeypatch.setattr(module,'connect_shooju',lambda:object())
    def retrieve(key,start,end,**kwargs):
        spec=sources.source_contract(repo_root)['sources'][key]
        response=npd_response(spec,pd.date_range(start,end,freq='MS'))
        if field=='hanz': response['values']=response['values'].iloc[2:]
        return response
    monkeypatch.setattr(sources,'get_registered_series',retrieve)
    assert module.activate(True,[reference])==0
    assert registry_path.read_bytes().startswith(original)
    _,registry=load_config(tmp_path/'config/model.yaml',registry_path)
    assert registry.at[key,'asset']==''
    assert registry.at[key,'sid'].endswith(field+'\\gas-saleable')
    assert module.activate(True,[reference])==0  # Recheck must not append a duplicate.
    assert len(registry_path.read_bytes().splitlines())==len(original.splitlines())+1
    with pytest.raises(ValueError,match='within January-December 2024'):
        module.activate(False,[reference],start='2026-01',end='2026-09')


def test_hanz_structural_zeros_are_specific_and_preserve_source_observations(repo_root):
    contract=sources.source_contract(repo_root)
    spec=contract['sources']['hanz_net_actual']
    dates=pd.date_range('2024-01-01','2024-12-01',freq='MS')
    response=npd_response(spec,dates)
    response['values']=response['values'].iloc[2:]
    original=response['values'].copy()
    result=sources.verify_response('hanz_net_actual',response,spec,dates,'monthly')
    assert result.iloc[:3].eq(0).all()
    pd.testing.assert_series_equal(response['values'],original)
    response['values'].iloc[0]=.1
    with pytest.raises(ValueError,match='contradicts'):
        sources.verify_response('hanz_net_actual',response,spec,dates,'monthly')
    other=contract['sources']['fram_net_actual']
    response=npd_response(other,dates)
    response['values']=response['values'].iloc[2:]
    with pytest.raises(ValueError,match='missing 2024-01-01, 2024-02-01'):
        sources.verify_response('fram_net_actual',response,other,dates,'monthly')
