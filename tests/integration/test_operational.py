"""Synthetic end-to-end workflow and fail-closed operational boundaries."""
from copy import deepcopy
import hashlib
import json
import shutil

import pandas as pd
from openpyxl import load_workbook
import pytest
import yaml

from norway_gas_routing import operational as op
from norway_gas_routing.routing import network_definition
from norway_gas_routing.reporting.operational import excel_value


@pytest.fixture
def workspace(repo_root, tmp_path):
    # Repository-local pytest basetemp. Copy only task source/configuration,
    # never retained raw data; Git identity uses this real checkout's identity.
    for directory in ['config', 'src/norway_gas_routing']:
        shutil.copytree(repo_root/directory, tmp_path/directory)
    shutil.copytree(repo_root/'evidence/inputs/routing_sources_b1', tmp_path/'evidence/inputs/routing_sources_b1')
    for directory in ['production_forecast_lineage','reconciled_supply_f22']:
        shutil.copytree(repo_root/'evidence/inputs'/directory,tmp_path/'evidence/inputs'/directory)
    (tmp_path/'scripts/run').mkdir(parents=True)
    shutil.copyfile(repo_root/'scripts/run/run_routing.py', tmp_path/'scripts/run/run_routing.py')
    return tmp_path


def configure(root, **changes):
    path = root/'config/run.yaml'
    request = yaml.safe_load(path.read_text(encoding='utf-8-sig'))
    request.update(mode='historical_reconstruction',capacity_scenario='baseline',origin=None,
                   input_source='synthetic_demo',start='2024-01',end='2024-02')
    request.update(changes)
    path.write_text(yaml.safe_dump(request), encoding='utf-8')
    return request


def test_complete_solve_excel_provenance_and_nonoverwrite(workspace):
    configure(workspace)
    first, code = op.execute_run(workspace)
    assert code == 0
    manifest = json.loads((first/'provenance/run_manifest.json').read_text())
    assert manifest['status'] == 'OPTIMAL'
    assert manifest['settings']['price_enabled'] is False
    assert manifest['git']['commit']
    assert manifest['evaluation']['holdout_scored'] is False
    assert 'SYNTHETIC' in manifest['analysis_type']
    for name, digest in manifest['files_sha256'].items():
        assert hashlib.sha256((first/name).read_bytes()).hexdigest() == digest
    workbook = load_workbook(first/'results.xlsx', data_only=False)
    assert workbook.sheetnames == ['Read me', 'Monthly flows', 'Comparisons', 'Capacities',
        'Network flows', 'Preferences', 'Input sources', 'Settings', 'Diagnostics']
    assert workbook['Monthly flows'].freeze_panes == 'B4'
    assert len(workbook['Monthly flows']._charts) == 1
    assert workbook['Comparisons']['E4'].value == 'Unavailable'
    assert workbook['Comparisons']['F4'].value == 'Unavailable'
    assert not any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row)
    csv = pd.read_csv(first/'results.csv')
    for index, row in csv.iterrows():
        assert workbook['Monthly flows'].cell(index+4, 9).value == pytest.approx(row.total_exports_bcm)
        assert sum(row[d+'_volume_bcm'] for d in op.DESTINATIONS) == pytest.approx(row.total_exports_bcm)
        assert row.total_exports_bcm+row.unrouted_supply_bcm == pytest.approx(row.pipeline_supply_proxy_bcm)
    inputs = pd.read_csv(first/'provenance/inputs/supply.csv')
    assert inputs.iloc[0].KARSTO_FEED == 1.5
    checks = pd.read_csv(first/'details/solver_checks.csv')
    assert checks.max_capacity_violation_bcm.max() < 1e-7
    assert checks.max_mass_balance_residual_bcm.max() < 1e-7
    assert (first/'details/monthly_flows.svg').is_file()
    assert 'results.xlsx' in (first/'report.html').read_text(encoding='utf-8')
    workbook.close()
    before = op.digest(first/'results.xlsx')
    second, code = op.execute_run(workspace, validate_only=True)
    assert code == 0 and first != second and op.digest(first/'results.xlsx') == before


def test_validate_only_never_builds_lp_or_retrieves(workspace, monkeypatch):
    configure(workspace, end='2024-01')
    def forbidden(*args, **kwargs):
        pytest.fail('Validation-only attempted optimisation')
    monkeypatch.setattr(op, 'build_month', forbidden)
    monkeypatch.setattr(op, 'available_solver', forbidden)
    monkeypatch.setattr(op, 'optimise', forbidden)
    output, code = op.execute_run(workspace, validate_only=True)
    assert code == 0
    manifest = json.loads((output/'provenance/run_manifest.json').read_text())
    assert manifest['status'] == 'VALIDATED_ONLY' and manifest['solver'] is None
    assert json.loads((output/'provenance/solver_results.json').read_text()) == {}
    book = load_workbook(output/'results.xlsx', read_only=True)
    assert book['Monthly flows']['B4'].value == 'Unavailable'
    book.close()


@pytest.mark.parametrize('changes,message', [
    ({'mode': 'retrospective_forecast', 'origin': '2023-12-31T23:59:59Z'}, 'complete origin bundle'),
    ({'mode': 'latest_forecast'}, 'Forecast mode unavailable'),
])
def test_blocked_modes_have_actionable_packages(workspace, changes, message):
    configure(workspace, **changes)
    output, code = op.execute_run(workspace, validate_only=True)
    assert code == 2
    manifest = json.loads((output/'provenance/run_manifest.json').read_text())
    assert manifest['status'] == 'BLOCKED_INPUTS'
    assert message in manifest['blocker']
    assert manifest['inputs'] == []
    assert not list((output/'provenance/inputs').iterdir())
    assert (output/'results.xlsx').exists()


@pytest.mark.parametrize('changes,message', [
    ({'start': '2026-01', 'end': '2026-01'}, 'locked'),
    ({'start': '2030-01', 'end': '2030-01'}, '2025'),
    ({'input_source': 'frozen'}, 'cached/frozen'),
    ({'price_enabled': True}, 'Experimental price'),
    ({'capacity_scenario': 'kollsnes-156'}, 'baseline capacity'),
    ({'evaluation': 'config/alternate.yaml'}, 'cannot be replaced'),
    ({'model': '../outside.yaml'}, 'leaves the authorised'),
    ({'origin': '2023-12-31'}, 'origin: null'),
    ({'end': '2023-12'}, 'start must'),
    ({'mode': []}, 'Unsupported mode'),
    ({'input_source': {}}, 'Unsupported input_source'),
])
def test_config_rejects_unsupported_requests_before_preparation(workspace, changes, message):
    configure(workspace, **changes)
    with pytest.raises(ValueError, match=message):
        op.execute_run(workspace, validate_only=True)
    assert not (workspace/'outputs').exists()


def test_solver_failure_still_produces_traceable_unsolved_package(workspace):
    configure(workspace, solver='nonexistent_solver')
    output, code = op.execute_run(workspace)
    assert code == 2
    manifest = json.loads((output/'provenance/run_manifest.json').read_text())
    assert manifest['status'] == 'FAILED_OPTIMISATION'
    assert manifest['blocker']
    assert pd.read_csv(output/'results.csv').total_exports_bcm.isna().all()
    assert (output/'results.xlsx').is_file()


def test_input_boundary_no_gaps_or_negative_supply(workspace):
    configure(workspace)
    request, model, _, _ = op.load_run_config(workspace, 'config/run.yaml')
    inputs = op.synthetic_inputs(model)
    network = network_definition(model['routing'])
    broken = deepcopy(inputs)
    broken['supply'].iloc[0, 0] = float('nan')
    with pytest.raises(ValueError, match='Nonfinite'):
        op.validate_prepared_inputs(broken, network, request)
    broken = deepcopy(inputs)
    broken['capacities'] = broken['capacities'].iloc[1:]
    with pytest.raises(ValueError, match='exact unique'):
        op.validate_prepared_inputs(broken, network, request)
    broken = deepcopy(inputs)
    broken['training_observations'] = broken['training_observations'].iloc[:0]
    with pytest.raises(ValueError, match='nonempty unique earlier'):
        op.validate_prepared_inputs(broken, network, request)
    assert excel_value('=HYPERLINK("malicious")').startswith("'")


def test_single_month_future_adapter_shape(workspace):
    configure(workspace, start='2025-01', end='2025-01')
    request, model, _, _ = op.load_run_config(workspace, 'config/run.yaml')
    inputs = op.synthetic_inputs(model)
    op.validate_prepared_inputs(inputs, network_definition(model['routing']), request)
    assert len(inputs['months']) == 1


def test_malformed_yaml_is_an_actionable_configuration_error(workspace):
    (workspace/'config/run.yaml').write_text('model: [broken', encoding='utf-8')
    with pytest.raises(ValueError, match='Invalid run YAML'):
        op.execute_run(workspace)
    assert not (workspace/'outputs').exists()
