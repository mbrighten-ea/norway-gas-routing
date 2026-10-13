"""Retained forecast replay/presentation tests; no service calls or new forecasts."""
import importlib.util
import json
import sys
from pathlib import Path
from openpyxl import load_workbook
import pandas as pd

ROOT=Path(__file__).resolve().parents[2]
SUMMARY='outputs/forward_scenarios/20261013_f26'


def command(name, directory='run'):
    spec=importlib.util.spec_from_file_location(name,ROOT/'scripts'/directory/(name+'.py'))
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    return module


def test_strict_released_snapshots_replay_without_mutating_package():
    module=command('verify_forward_scenarios','verify')
    path=ROOT/SUMMARY/'provenance/validation.json';original=path.read_bytes()
    module.main([SUMMARY,'--replay'])
    assert path.read_bytes()==original


def test_one_workbook_contains_four_scenarios_accounting_and_capacity(tmp_path):
    module=command('run_forward_scenarios')
    manifest=json.loads((ROOT/SUMMARY/'provenance/manifest.json').read_text())
    # Use retained completed results as immutable report input, not solver inputs.
    import shutil
    (tmp_path/'details').mkdir()
    for file in ['results.csv','details/annual_allocations.csv','details/scenario_differences.csv']:
        shutil.copyfile(ROOT/SUMMARY/file,tmp_path/file)
    module.write_summary_workbook(tmp_path,manifest)
    book=load_workbook(tmp_path/'results.xlsx',data_only=True)
    assert book.sheetnames[:3]==['Monthly scenarios','Annual scenarios','Scenario differences']
    assert {'Accounting','Source nodes','Capacity utilisation','Read me'}.issubset(book.sheetnames)
    sheet=book['Monthly scenarios'];assert sheet.max_row==155
    headers=[c.value for c in sheet[3]]
    rows=list(sheet.iter_rows(min_row=4,values_only=True))
    for label in 'ABCD':
        subset=[r for r in rows if r[headers.index('scenario')]==label]
        assert len(subset)==38
        assert [r[0].strftime('%Y-%m') for r in subset]==list(pd.period_range('2026-11','2029-12',freq='M').astype(str))
    limits=book['Capacity utilisation'];columns=[c.value for c in limits[3]]
    for row in limits.iter_rows(min_row=4,values_only=True):
        value=row[columns.index('utilisation_pct')]
        assert value=='Unavailable' or 0<=value<=100+1e-7
    book.close()


def test_normal_command_dispatches_configured_comparison_and_validation(monkeypatch):
    cli=command('run_routing');batch=command('run_forward_scenarios')
    calls=[]
    monkeypatch.setitem(sys.modules,'run_forward_scenarios',batch)
    monkeypatch.setattr(cli,'load_run_config',lambda *a:({'compare_scenarios':True},None,None,None))
    monkeypatch.setattr(batch,'main',lambda config,**kwargs:calls.append((config,kwargs)) or 0)
    monkeypatch.setattr(cli,'execute_run',lambda *a,**k: (_ for _ in ()).throw(AssertionError('Individual solve must not run')))
    assert cli.main(['--config','config/run.yaml','--validate-only','--solver','HiGHS'])==0
    assert calls==[('config/run.yaml',{'validate_only':True,'solver_name':'HiGHS'})]


def test_four_case_validation_writes_one_summary_without_solving(tmp_path,monkeypatch):
    module=command('run_forward_scenarios');calls=[]
    monkeypatch.setattr(module,'ROOT',tmp_path)
    settings={'physical_input_contract':'core_f24','origin':'2026-10-13T13:49:52+00:00',
        'start':'2026-11','end':'2029-12','solver':'HiGHS'}
    monkeypatch.setattr(module,'load_run_config',lambda *a:(settings,None,None,None))
    def validate(root,config,**kwargs):
        import yaml
        request=yaml.safe_load((root/config).read_text());calls.append((request,kwargs))
        path=root/'outputs/routing_runs'/str(len(calls))
        (path/'provenance/raw').mkdir(parents=True)
        (path/'provenance/run_manifest.json').write_text(json.dumps({'status':'VALIDATED_ONLY','blocker':None}))
        return path,0
    monkeypatch.setattr(module,'execute_run',validate)
    assert module.main('config/run.yaml',validate_only=True)==0
    assert [(r['heidrun_treatment'],r['capacity_scenario']) for r,k in calls]==[
        ('export_eligible','core_f24_156'),('domestic_bound','core_f24_156'),
        ('export_eligible','core_f24_proportional_143'),('domestic_bound','core_f24_proportional_143')]
    assert all(k['validate_only'] and r['origin']==settings['origin'] for r,k in calls)
    summary=next((tmp_path/'outputs/forward_scenarios').iterdir())
    assert json.loads((summary/'provenance/manifest.json').read_text())['status']=='VALIDATED_ONLY'
    book=load_workbook(summary/'results.xlsx',data_only=True)
    assert book['Monthly scenarios'].max_row==3
    book.close()
